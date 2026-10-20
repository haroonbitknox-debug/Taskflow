import csv
import re
import time
import uuid
import zipfile
from collections.abc import Iterator
from datetime import date, datetime, time as datetime_time
from pathlib import Path, PurePosixPath
from typing import Any

import pandas as pd
from fastapi import UploadFile
from openpyxl import load_workbook
from sqlalchemy import String, func, insert, literal, or_, select, union_all
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models import FileIndexEntry, FileProcessingStatus, UploadedFile, User
from app.schemas.files import FileMetadata, FilePage, FileSearchHit, FileSearchPage

MIME_TYPES = {
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".xls": "application/vnd.ms-excel",
    ".csv": "text/csv",
}
ALLOWED_MIME_TYPES = {
    ".xlsx": {
        MIME_TYPES[".xlsx"],
        "application/zip",
        "application/octet-stream",
    },
    ".xls": {MIME_TYPES[".xls"], "application/octet-stream"},
    ".csv": {MIME_TYPES[".csv"], "application/csv", "application/octet-stream"},
}
UPLOADS_BASE = Path(__file__).resolve().parents[2]


class FileServiceError(Exception):
    status_code = 400


class FileValidationError(FileServiceError):
    status_code = 422


class FileNotFoundError(FileServiceError):
    status_code = 404


class FileStorageError(FileServiceError):
    status_code = 500


def storage_root() -> Path:
    configured = Path(settings.upload_storage_dir)
    if not configured.is_absolute():
        configured = UPLOADS_BASE / configured
    return configured.resolve()


def _safe_storage_path(file: UploadedFile) -> Path:
    root = storage_root()
    relative = PurePosixPath(file.relative_path)
    if relative.is_absolute() or any(part in {"", ".", ".."} for part in relative.parts):
        raise FileStorageError("Stored file path is invalid.")
    candidate = (root / Path(*relative.parts)).resolve()
    if root not in candidate.parents:
        raise FileStorageError("Stored file path is outside the upload directory.")
    return candidate


def _clean_filename(filename: str | None) -> str:
    name = (filename or "").replace("\\", "/").split("/")[-1].strip()
    if (
        not name
        or name in {".", ".."}
        or len(name) > 255
        or any(ord(character) < 32 or ord(character) == 127 for character in name)
    ):
        raise FileValidationError("Provide a valid filename no longer than 255 characters.")
    return name


def _remove_storage_file(path: Path) -> None:
    try:
        path.unlink(missing_ok=True)
    except OSError as exc:
        raise FileStorageError("Could not clean up the stored file.") from exc


def _validate_mime(extension: str, content_type: str | None) -> None:
    mime = (content_type or "").split(";", 1)[0].strip().casefold()
    if mime and mime not in ALLOWED_MIME_TYPES[extension]:
        raise FileValidationError("The uploaded content type does not match the file extension.")


def _validate_signature(path: Path, extension: str) -> None:
    with path.open("rb") as source:
        signature = source.read(8)
    if extension == ".xlsx":
        if not zipfile.is_zipfile(path):
            raise FileValidationError("This is not a valid XLSX workbook.")
        try:
            with zipfile.ZipFile(path) as archive:
                members = archive.namelist()
                if "xl/workbook.xml" not in members:
                    raise FileValidationError("The XLSX workbook is missing its workbook data.")
                if any(name.casefold().endswith("vbaproject.bin") for name in members):
                    raise FileValidationError("Macro-enabled workbooks are not supported.")
                if sum(item.file_size for item in archive.infolist()) > settings.file_max_extracted_text_bytes * 5:
                    raise FileValidationError("The workbook expands beyond the allowed processing size.")
        except zipfile.BadZipFile as exc:
            raise FileValidationError("This is not a valid XLSX workbook.") from exc
    elif extension == ".xls":
        if signature != bytes.fromhex("D0CF11E0A1B11AE1"):
            raise FileValidationError("This is not a valid legacy XLS workbook.")
    elif b"\x00" in signature:
        raise FileValidationError("CSV files must contain plain text data.")


async def _write_upload(upload: UploadFile, path: Path) -> int:
    size = 0
    try:
        with path.open("xb") as destination:
            while chunk := await upload.read(64 * 1024):
                size += len(chunk)
                if size > settings.file_max_size_bytes:
                    raise FileValidationError(
                        f"File exceeds the {settings.file_max_size_bytes // (1024 * 1024)} MB upload limit."
                    )
                destination.write(chunk)
    except Exception:
        _remove_storage_file(path)
        raise
    if size == 0:
        _remove_storage_file(path)
        raise FileValidationError("The uploaded file is empty.")
    return size


def _column_label(index: int) -> str:
    label = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        label = chr(65 + remainder) + label
    return label or "A"


def _contains_pattern(value: str) -> str:
    escaped = value.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
    return f"%{escaped}%"


def _headers(values: list[Any], width: int) -> list[str]:
    result: list[str] = []
    used: dict[str, int] = {}
    for index in range(width):
        value = _cell_text(values[index] if index < len(values) else None)
        base = value[:64] if value else f"Column {_column_label(index + 1)}"
        count = used.get(base.casefold(), 0) + 1
        used[base.casefold()] = count
        name = base if count == 1 else f"{base[:58]} ({count})"
        result.append(name[:64])
    return result


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (datetime, date, datetime_time)):
        return value.isoformat()
    return str(value).strip()


def _cell_rows(path: Path, extension: str) -> Iterator[tuple[str, int, str, str]]:
    started = time.monotonic()
    extracted_bytes = 0
    sheets = 0

    def check_time() -> None:
        if time.monotonic() - started > settings.file_max_processing_seconds:
            raise FileValidationError("Spreadsheet processing exceeded the configured time limit.")

    def yield_rows(sheet_name: str, rows: Iterator[list[Any]]) -> Iterator[tuple[str, int, str, str]]:
        nonlocal extracted_bytes
        header_labels: list[str] | None = None
        width = 0
        for row_number, row in enumerate(rows, start=1):
            check_time()
            if row_number > settings.file_max_rows + 1:
                raise FileValidationError(
                    f"Worksheet exceeds the {settings.file_max_rows} row limit."
                )
            if len(row) > settings.file_max_columns:
                raise FileValidationError(
                    f"Worksheet exceeds the {settings.file_max_columns} column limit."
                )
            width = max(width, len(row))
            if row_number == 1:
                header_labels = _headers(row, width)
                continue
            if header_labels is None or len(header_labels) < width:
                header_labels = _headers([], width)
            for index, value in enumerate(row):
                text = _cell_text(value)
                if not text:
                    continue
                if "\x00" in text:
                    raise FileValidationError("Spreadsheet values must not contain binary null bytes.")
                content = f"{header_labels[index]}: {text}"
                extracted_bytes += len(content.encode("utf-8"))
                if extracted_bytes > settings.file_max_extracted_text_bytes:
                    raise FileValidationError("Extracted spreadsheet text exceeds the configured limit.")
                yield sheet_name[:128], row_number, header_labels[index], content
        check_time()

    if extension == ".csv":
        try:
            with path.open("r", encoding="utf-8-sig", newline="") as source:
                rows = csv.reader(source, strict=True)
                sheets = 1
                yield from yield_rows("CSV", rows)
        except (UnicodeDecodeError, csv.Error) as exc:
            raise FileValidationError("CSV must be valid UTF-8 text with well-formed rows.") from exc
        return

    if extension == ".xlsx":
        try:
            workbook = load_workbook(
                path,
                read_only=True,
                data_only=True,
                keep_links=False,
            )
        except Exception as exc:
            raise FileValidationError("The Excel workbook could not be read.") from exc
        try:
            sheets = len(workbook.sheetnames)
            if not sheets:
                raise FileValidationError("The workbook contains no worksheets.")
            if sheets > settings.file_max_worksheets:
                raise FileValidationError(
                    f"Workbook exceeds the {settings.file_max_worksheets} worksheet limit."
                )
            for worksheet in workbook.worksheets:
                rows = (list(row) for row in worksheet.iter_rows(values_only=True))
                yield from yield_rows(worksheet.title or "Worksheet", rows)
        finally:
            workbook.close()
        return

    try:
        workbook = pd.ExcelFile(path, engine="xlrd")
    except Exception as exc:
        raise FileValidationError("The legacy XLS workbook could not be read.") from exc
    try:
        sheets = len(workbook.sheet_names)
        if not sheets:
            raise FileValidationError("The workbook contains no worksheets.")
        if sheets > settings.file_max_worksheets:
            raise FileValidationError(
                f"Workbook exceeds the {settings.file_max_worksheets} worksheet limit."
            )
        for sheet_name in workbook.sheet_names:
            check_time()
            frame = workbook.parse(
                sheet_name=sheet_name,
                header=None,
                nrows=settings.file_max_rows + 2,
                dtype=object,
            )
            if len(frame.index) > settings.file_max_rows + 1:
                raise FileValidationError(
                    f"Worksheet exceeds the {settings.file_max_rows} row limit."
                )
            if len(frame.columns) > settings.file_max_columns:
                raise FileValidationError(
                    f"Worksheet exceeds the {settings.file_max_columns} column limit."
                )
            yield from yield_rows(
                sheet_name,
                (list(values) for values in frame.itertuples(index=False, name=None)),
            )
    finally:
        workbook.close()


def _metadata(db: Session, file: UploadedFile) -> FileMetadata:
    count = db.scalar(
        select(func.count(FileIndexEntry.id)).where(
            FileIndexEntry.uploaded_file_id == file.id,
            FileIndexEntry.user_id == file.user_id,
        )
    ) or 0
    return FileMetadata(
        id=file.id,
        original_filename=file.original_filename,
        extension=file.extension,
        mime_type=file.mime_type,
        file_size=file.file_size,
        created_at=file.created_at,
        processing_status=file.processing_status.value,
        indexing_status=file.processing_status.value,
        indexed_entry_count=count,
        processing_error=file.processing_error,
    )


def _run_indexing(db: Session, file: UploadedFile, path: Path) -> None:
    file.processing_status = FileProcessingStatus.PROCESSING
    file.processing_error = None
    try:
        db.commit()
    except SQLAlchemyError as exc:
        db.rollback()
        current = db.get(UploadedFile, file.id)
        if current is not None:
            current.processing_status = FileProcessingStatus.FAILED
            current.processing_error = "Indexing failed. You can safely retry processing."
            try:
                db.commit()
            except SQLAlchemyError as status_error:
                db.rollback()
                raise FileStorageError(
                    "Could not start indexing or save its failure status."
                ) from status_error
        raise FileStorageError("Could not start spreadsheet indexing.") from exc
    batch: list[dict[str, Any]] = []
    try:
        for sheet, row_number, column, content in _cell_rows(path, file.extension):
            batch.append({
                "uploaded_file_id": file.id,
                "user_id": file.user_id,
                "worksheet_name": sheet,
                "row_number": row_number,
                "column_name": column,
                "searchable_content": content,
                "normalized_content": re.sub(
                    r"[\W_]+", " ", content.casefold(), flags=re.UNICODE
                ).strip()[:512],
            })
            if len(batch) >= 500:
                db.execute(insert(FileIndexEntry), batch)
                batch.clear()
        if batch:
            db.execute(insert(FileIndexEntry), batch)
        file.processing_status = FileProcessingStatus.READY
        file.processing_error = None
        db.commit()
    except FileValidationError:
        db.rollback()
        current = db.get(UploadedFile, file.id)
        if current is not None:
            db.delete(current)
            try:
                db.commit()
            except SQLAlchemyError as exc:
                db.rollback()
                raise FileStorageError(
                    "Could not remove invalid file metadata after indexing failed."
                ) from exc
        _remove_storage_file(path)
        raise
    except Exception as exc:
        db.rollback()
        current = db.get(UploadedFile, file.id)
        if current is None:
            raise FileStorageError("File processing failed and its metadata could not be restored.") from exc
        current.processing_status = FileProcessingStatus.FAILED
        current.processing_error = "Indexing failed. You can safely retry processing."
        try:
            db.commit()
        except SQLAlchemyError as db_error:
            db.rollback()
            raise FileStorageError("File indexing failed and processing status could not be saved.") from db_error


async def upload_file(db: Session, user: User, upload: UploadFile) -> FileMetadata:
    original_name = _clean_filename(upload.filename)
    extension = Path(original_name).suffix.casefold()
    if extension not in MIME_TYPES:
        raise FileValidationError("Only .xlsx, .xls, and .csv files are supported.")
    _validate_mime(extension, upload.content_type)

    root = storage_root()
    user_directory = root / str(user.id)
    try:
        user_directory.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        raise FileStorageError("Could not create the upload storage directory.") from exc
    stored_name = f"{uuid.uuid4().hex}{extension}"
    destination = user_directory / stored_name
    relative_path = f"{user.id}/{stored_name}"
    try:
        size = await _write_upload(upload, destination)
    except FileServiceError:
        raise
    except OSError as exc:
        raise FileStorageError("Could not write the uploaded file.") from exc
    try:
        _validate_signature(destination, extension)
        record = UploadedFile(
            user_id=user.id,
            original_filename=original_name,
            stored_filename=stored_name,
            relative_path=relative_path,
            extension=extension,
            mime_type=MIME_TYPES[extension],
            file_size=size,
            processing_status=FileProcessingStatus.PENDING,
        )
        db.add(record)
        db.commit()
        db.refresh(record)
    except FileServiceError:
        _remove_storage_file(destination)
        raise
    except SQLAlchemyError as exc:
        db.rollback()
        _remove_storage_file(destination)
        raise FileStorageError("Could not save uploaded file metadata.") from exc

    try:
        _run_indexing(db, record, destination)
    except FileValidationError:
        raise
    db.refresh(record)
    return _metadata(db, record)


def list_files(
    db: Session,
    user_id: int,
    page: int,
    page_size: int,
    extension: str | None = None,
) -> FilePage:
    statement = select(UploadedFile).where(UploadedFile.user_id == user_id)
    if extension:
        statement = statement.where(UploadedFile.extension == extension.casefold())
    total = db.scalar(select(func.count()).select_from(statement.subquery())) or 0
    files = db.scalars(
        statement.order_by(UploadedFile.created_at.desc(), UploadedFile.id.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
    ).all()
    return FilePage(
        items=[_metadata(db, file) for file in files],
        total=total,
        page=page,
        page_size=page_size,
    )


def get_file(db: Session, file_id: int, user_id: int) -> UploadedFile:
    file = db.scalar(
        select(UploadedFile).where(
            UploadedFile.id == file_id,
            UploadedFile.user_id == user_id,
        )
    )
    if file is None:
        raise FileNotFoundError("File not found.")
    return file


def get_file_metadata(db: Session, file_id: int, user_id: int) -> FileMetadata:
    return _metadata(db, get_file(db, file_id, user_id))


def download_path(db: Session, file_id: int, user_id: int) -> tuple[Path, UploadedFile]:
    file = get_file(db, file_id, user_id)
    path = _safe_storage_path(file)
    if not path.is_file():
        raise FileNotFoundError("Stored file is no longer available.")
    return path, file


def delete_file(db: Session, file_id: int, user_id: int) -> None:
    file = get_file(db, file_id, user_id)
    path = _safe_storage_path(file)
    staged_path = path.with_name(f".{path.name}.{uuid.uuid4().hex}.deleting")
    if path.exists():
        try:
            path.replace(staged_path)
        except OSError as exc:
            raise FileStorageError("Could not stage the stored file for deletion.") from exc
    db.delete(file)
    try:
        db.commit()
    except SQLAlchemyError as exc:
        db.rollback()
        if staged_path.exists():
            try:
                staged_path.replace(path)
            except OSError as restore_error:
                raise FileStorageError(
                    "File metadata deletion failed and the stored file could not be restored."
                ) from restore_error
        raise FileStorageError("Could not delete file metadata.") from exc
    if staged_path.exists():
        try:
            staged_path.unlink()
        except OSError as exc:
            raise FileStorageError(
                "File metadata was deleted, but stored file cleanup failed. Contact an administrator."
            ) from exc


def retry_indexing(db: Session, file_id: int, user_id: int) -> FileMetadata:
    file = get_file(db, file_id, user_id)
    if file.processing_status != FileProcessingStatus.FAILED:
        raise FileValidationError("Only failed files can be retried.")
    path = _safe_storage_path(file)
    if not path.is_file():
        raise FileNotFoundError("Stored file is no longer available.")
    db.execute(
        FileIndexEntry.__table__.delete().where(
            FileIndexEntry.uploaded_file_id == file.id,
            FileIndexEntry.user_id == user_id,
        )
    )
    file.processing_status = FileProcessingStatus.PENDING
    file.processing_error = None
    db.commit()
    _run_indexing(db, file, path)
    db.refresh(file)
    return _metadata(db, file)


def search_files(
    db: Session,
    user_id: int,
    query: str,
    page: int,
    page_size: int,
    file_id: int | None = None,
    extension: str | None = None,
) -> FileSearchPage:
    search = _contains_pattern(query.strip())
    normalized_query = re.sub(
        r"[\W_]+", " ", query.casefold(), flags=re.UNICODE
    ).strip()
    normalized_search = _contains_pattern(normalized_query)
    file_filter = [
        UploadedFile.user_id == user_id,
    ]
    if file_id is not None:
        file_filter.append(UploadedFile.id == file_id)
    if extension:
        file_filter.append(UploadedFile.extension == extension.casefold())

    filename_hits = select(
        UploadedFile.id.label("file_id"),
        UploadedFile.original_filename.label("original_filename"),
        UploadedFile.extension.label("extension"),
        literal(None, type_=String(128)).label("worksheet_name"),
        literal(None).label("row_number"),
        literal(None, type_=String(64)).label("column_name"),
        literal("filename").label("match_type"),
        UploadedFile.original_filename.label("excerpt"),
        UploadedFile.created_at.label("created_at"),
    ).where(*file_filter, UploadedFile.original_filename.ilike(search, escape="\\"))

    cell_filter = [
        FileIndexEntry.user_id == user_id,
        UploadedFile.user_id == user_id,
        or_(
            FileIndexEntry.worksheet_name.ilike(search, escape="\\"),
            FileIndexEntry.normalized_content.like(normalized_search, escape="\\")
            if normalized_query
            else FileIndexEntry.worksheet_name.ilike(search, escape="\\"),
        ),
    ]
    if file_id is not None:
        cell_filter.append(UploadedFile.id == file_id)
    if extension:
        cell_filter.append(UploadedFile.extension == extension.casefold())
    content_hits = (
        select(
            UploadedFile.id.label("file_id"),
            UploadedFile.original_filename.label("original_filename"),
            UploadedFile.extension.label("extension"),
            FileIndexEntry.worksheet_name.label("worksheet_name"),
            FileIndexEntry.row_number.label("row_number"),
            FileIndexEntry.column_name.label("column_name"),
            literal("content").label("match_type"),
            FileIndexEntry.searchable_content.label("excerpt"),
            UploadedFile.created_at.label("created_at"),
        )
        .join(FileIndexEntry, FileIndexEntry.uploaded_file_id == UploadedFile.id)
        .where(*cell_filter)
    )
    hits = union_all(filename_hits, content_hits).subquery()
    total = db.scalar(select(func.count()).select_from(hits)) or 0
    rows = db.execute(
        select(hits)
        .order_by(
            hits.c.created_at.desc(),
            hits.c.file_id.desc(),
            hits.c.row_number.asc(),
            hits.c.column_name.asc(),
            hits.c.match_type.asc(),
        )
        .offset((page - 1) * page_size)
        .limit(page_size)
    ).all()
    items = []
    for row in rows:
        data = row._mapping
        match_type = data["match_type"]
        if match_type == "content":
            match_type = "worksheet" if query.casefold() in data["worksheet_name"].casefold() else "content"
        excerpt = re.sub(r"\s+", " ", data["excerpt"]).strip()
        items.append(
            FileSearchHit(
                file_id=data["file_id"],
                original_filename=data["original_filename"],
                extension=data["extension"],
                worksheet_name=data["worksheet_name"],
                row_number=data["row_number"],
                column_name=data["column_name"],
                match_type=match_type,
                excerpt=excerpt[:240],
                created_at=data["created_at"],
            )
        )
    return FileSearchPage(items=items, total=total, page=page, page_size=page_size)
